from dataclasses import dataclass
from pathlib import Path
from collections import defaultdict
from typing import Dict, List, Optional, Tuple
import math
import random
import re

from tqdm import tqdm

from openpyxl import load_workbook
from PIL import Image

import torch
from torch.utils.data import Dataset
from torchvision import transforms

from pathlib import Path
# from dataloader import build_records


@dataclass(frozen=True)
class SampleRecord:
    image_path: str
    family: str
    condition_id: str
    ply: int
    impactor: str
    j_per_mm: float
    angle: str
    impact_energy_j: float
    excel_path: str


def normalize_value(value) -> str:
    """
    将 Excel 或文件名中的编号统一成字符串。

    例如：
    1.0 -> "1"
    "6T" -> "6t"
    """
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip().lower()


def normalize_header(value) -> str:
    if value is None:
        return ""

    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def parse_image_name(image_path: Path) -> Optional[Tuple[str, str]]:
    """
    解析图片名。

    支持：

    c16-13.jpg
    c24-12t.jpg
    Q24-1astm.jpg

    返回：

    family: c16 / c24 / q24
    condition_id: 13 / 12t / 1astm
    """
    stem = image_path.stem.strip().lower()

    match = re.fullmatch(
        r"(?P<family>[cq](?:8|16|24))-(?P<condition>.+)",
        stem,
    )

    if match is None:
        return None

    family = match.group("family")
    condition_id = normalize_value(match.group("condition"))

    return family, condition_id


def choose_excel_file(
    family: str,
    condition_id: str,
    condition_dir: Path,
) -> Tuple[Path, str]:
    """
    根据图片所属类型选择 Excel 文件。

    Q24-1astm.jpg 对应：
        impact condition_astm.xlsx

    普通图片对应：
        impact condition_c8.xlsx
        impact condition_c16.xlsx
        impact condition_c24.xlsx
        impact condition_q8.xlsx
        impact condition_q16.xlsx
        impact condition_q24.xlsx
    """
    if family == "q24" and condition_id.endswith("astm"):
        excel_path = condition_dir / "impact condition_astm.xlsx"
        normalized_condition_id = condition_id[:-4]
    else:
        excel_path = condition_dir / f"impact condition_{family}.xlsx"
        normalized_condition_id = condition_id

    return excel_path, normalized_condition_id


def read_condition_excel(excel_path: Path) -> Dict[str, dict]:
    """
    读取一个 Excel 条件表。

    返回格式：

    {
        "13": {
            "ply": 16,
            "j_per_mm": 3.35,
            "impactor": "HemiA",
            "angle": "counter crockwise",
        }
    }
    """
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
    )

    if "condition" in workbook.sheetnames:
        worksheet = workbook["condition"]
    else:
        worksheet = workbook.active

    rows = worksheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        raise ValueError(f"Empty Excel file: {excel_path}")

    header_map = {
        normalize_header(value): index
        for index, value in enumerate(header_row)
    }

    required_columns = [
        "ply",
        "no.",
        "impactor",
        "j/mm",
        "angle",
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in header_map
    ]

    if missing_columns:
        raise ValueError(
            f"Missing columns {missing_columns} in {excel_path}"
        )

    condition_table = {}

    for row in rows:
        if not row:
            continue

        condition_id = normalize_value(row[header_map["no."]])

        if not condition_id:
            continue

        ply_value = row[header_map["ply"]]
        j_per_mm_value = row[header_map["j/mm"]]

        if ply_value is None or j_per_mm_value is None:
            continue

        try:
            ply = int(float(ply_value))
            j_per_mm = float(j_per_mm_value)
        except (TypeError, ValueError):
            continue

        if not math.isfinite(j_per_mm):
            continue

        condition_table[condition_id] = {
            "ply": ply,
            "j_per_mm": j_per_mm,
            "impactor": str(
                row[header_map["impactor"]] or ""
            ).strip(),
            "angle": str(
                row[header_map["angle"]] or ""
            ).strip(),
        }

    workbook.close()

    return condition_table


def build_records(
    image_dir: str = "dataset/all_c_scans_new",
    condition_dir: str = "dataset/impact_conditions",
    thickness_per_ply_mm: float = 0.1875,
) -> Tuple[List[SampleRecord], List[str]]:
    """
    建立图片和标签之间的对应关系。

    返回：

    records:
        成功匹配的样本

    unmatched:
        没有 Excel 标签的图片路径
    """
    image_dir = Path(image_dir)
    condition_dir = Path(condition_dir)

    image_paths = sorted(
        list(image_dir.glob("*.jpg"))
        + list(image_dir.glob("*.JPG")),
        key=lambda path: path.name.lower(),
    )

    if not image_paths:
        raise FileNotFoundError(
            f"No JPG files found in {image_dir}"
        )

    excel_cache = {}
    records = []
    unmatched = []

    for image_path in tqdm(
        image_paths,
        desc="Indexing C-scan images",
        unit="image",
    ):
        parsed = parse_image_name(image_path)

        if parsed is None:
            unmatched.append(str(image_path))
            tqdm.write(
                f"[skip] Cannot parse image name: {image_path.name}"
            )
            continue

        family, condition_id = parsed

        excel_path, normalized_condition_id = choose_excel_file(
            family=family,
            condition_id=condition_id,
            condition_dir=condition_dir,
        )

        excel_key = str(excel_path)

        if excel_key not in excel_cache:
            try:
                excel_cache[excel_key] = read_condition_excel(
                    excel_path
                )
            except Exception as error:
                tqdm.write(
                    f"[skip] {image_path.name}: {error}"
                )
                unmatched.append(str(image_path))
                continue

        condition_table = excel_cache[excel_key]
        condition = condition_table.get(normalized_condition_id)

        if condition is None:
            unmatched.append(str(image_path))
            tqdm.write(
                "[skip] No matching Excel row: "
                f"{image_path.name} -> "
                f"{excel_path.name}, "
                f"no.={normalized_condition_id}"
            )
            continue

        impact_energy_j = (
            condition["j_per_mm"]
            * condition["ply"]
            * thickness_per_ply_mm
        )

        records.append(
            SampleRecord(
                image_path=str(image_path),
                family=family,
                condition_id=normalized_condition_id,
                ply=condition["ply"],
                impactor=condition["impactor"],
                j_per_mm=condition["j_per_mm"],
                angle=condition["angle"],
                impact_energy_j=float(impact_energy_j),
                excel_path=str(excel_path),
            )
        )

    print()
    print(f"Total images: {len(image_paths)}")
    print(f"Matched samples: {len(records)}")
    print(f"Unmatched images: {len(unmatched)}")

    return records, unmatched


def split_records(
    records: List[SampleRecord],
    validation_ratio: float = 0.2,
    seed: int = 42,
) -> Tuple[List[SampleRecord], List[SampleRecord]]:
    """
    将样本划分为训练集和验证集。

    按 family + condition_id 分组，避免同一个试件同时出现在
    train 和 validation 中。
    """
    if not 0.0 < validation_ratio < 1.0:
        raise ValueError(
            "validation_ratio must be between 0 and 1"
        )

    groups = defaultdict(list)

    for record in records:
        group_key = (
            f"{record.excel_path}:{record.family}:{record.condition_id}"
        )
        groups[group_key].append(record)

    group_keys = list(groups.keys())

    rng = random.Random(seed)
    rng.shuffle(group_keys)

    validation_group_count = max(
        1,
        int(round(len(group_keys) * validation_ratio)),
    )

    validation_groups = set(
        group_keys[:validation_group_count]
    )

    train_records = []
    validation_records = []

    for group_key, group_records in groups.items():
        if group_key in validation_groups:
            validation_records.extend(group_records)
        else:
            train_records.extend(group_records)

    train_records.sort(key=lambda record: record.image_path)
    validation_records.sort(
        key=lambda record: record.image_path
    )

    return train_records, validation_records


def build_transforms(
    image_size: int = 256,
):
    normalization = transforms.Normalize(
        mean=[0.485, 0.456, 0.406],
        std=[0.229, 0.224, 0.225],
    )

    return transforms.Compose(
        [
            transforms.Resize(
                (image_size, image_size)
            ),
            transforms.ToTensor(),
            normalization,
        ]
    )

class CScanDataset(Dataset):
    """
    Impact Energy 回归数据集。

    每次返回：

    {
        "image": Tensor[C, H, W],
        "target": Tensor[],
        "path": str,
        "family": str,
        "condition_id": str,
    }
    """

    def __init__(
        self,
        records: List[SampleRecord],
        transform=None,
        augment: bool = False,
    ):
        self.records = list(records)
        self.transform = transform
        self.augment = augment

    def __len__(self) -> int:
        if self.augment:
            return len(self.records) * 6
        return len(self.records)

    def __getitem__(self, index: int) -> dict:
        from torchvision.transforms import functional as TF
        if self.augment:
            record_index = index // 6
            augmentation_index = index % 6
        else:
            record_index = index
            augmentation_index = 0

        record = self.records[record_index]

        with Image.open(record.image_path) as image:
            image = image.convert("RGB")
            image = image.copy()

        if self.augment:

            if augmentation_index == 1:
                image = TF.rotate(image, 90)

            elif augmentation_index == 2:
                image = TF.rotate(image, 180)

            elif augmentation_index == 3:
                image = TF.rotate(image, 270)

            elif augmentation_index == 4:
                image = TF.hflip(image)

            elif augmentation_index == 5:
                image = TF.vflip(image)

        if self.transform is not None:
            image = self.transform(image)

        target = torch.tensor(
            record.impact_energy_j,
            dtype=torch.float32,
        )

        return {
            "image": image,
            "target": target,
            "path": record.image_path,
            "family": record.family,
            "condition_id": record.condition_id,
        }


def build_datasets(
    image_dir: str = "dataset/all_c_scans_new",
    condition_dir: str = "dataset/impact_conditions",
    image_size: int = 224,
    validation_ratio: float = 0.2,
    seed: int = 42,
    thickness_per_ply_mm: float = 0.1875,
):
    """
    一次性构造训练集和验证集。
    """
    records, unmatched = build_records(
        image_dir=image_dir,
        condition_dir=condition_dir,
        thickness_per_ply_mm=thickness_per_ply_mm,
    )

    if not records:
        raise RuntimeError(
            "No valid image-label pairs were found."
        )

    train_records, validation_records = split_records(
        records=records,
        validation_ratio=validation_ratio,
        seed=seed,
    )

    train_dataset = CScanDataset(
        records=train_records,
        transform=build_transforms(
            image_size=image_size,
            train=True,
            augment=True,
        ),
    )

    validation_dataset = CScanDataset(
        records=validation_records,
        transform=build_transforms(
            image_size=image_size,
            train=False,
        ),
        augment=False,
    )

    print()
    print(f"Train samples: {len(train_dataset)}")
    print(f"Validation samples: {len(validation_dataset)}")

    if unmatched:
        print()
        print("Unmatched files:")
        for path in unmatched:
            print(f"  - {path}")

    return train_dataset, validation_dataset

if __name__ == "__main__":


    records, _ = build_records()

    special = [
        r for r in records
        if "astm" not in Path(r.image_path).stem.lower()
        and Path(r.image_path).stem.lower().endswith("t")
    ]

    print("t samples:", len(special))

    for r in special:
        print(
            Path(r.image_path).name,
            r.impact_energy_j,
            r.family,
            r.condition_id,
        )